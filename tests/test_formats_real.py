"""Real binary fixtures and workers; no substituted Docling results."""

import io

from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table, TableRow, TableCell
from odf.text import P
from pptx import Presentation
import xlwt
import pytest

from dslm3.service import Application


def fixtures():
    book = xlwt.Workbook()
    sheet = book.add_sheet("Vega")
    sheet.write(0, 0, "città")
    sheet.write(1, 0, 42)
    xls = io.BytesIO()
    book.save(xls)
    doc = OpenDocumentSpreadsheet()
    table = Table(name="Vega")
    row = TableRow()
    cell = TableCell(valuetype="float", value="3", formula="of:=1+2")
    cell.addElement(P(text="3"))
    row.addElement(cell)
    table.addElement(row)
    doc.spreadsheet.addElement(table)
    ods = io.BytesIO()
    doc.save(ods)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Vega ricambi"
    slide.placeholders[1].text = "Priorità P1: presa in carico"
    pptx = io.BytesIO()
    presentation.save(pptx)
    return [
        ("data.csv", 'nome;valore\n"città";12\n'.encode(), "dataset_row"),
        ("data.tsv", b"name\tvalue\nvega\t12\n", "dataset_row"),
        ("data.json", b'{"vega":12}', "structured_document"),
        ("data.jsonl", b'{"vega":12}\n{"vega":13}\n', "dataset_record"),
        ("data.yaml", b"date: 2026-01-01\nvega: 12\n", "structured_document"),
        (
            "data.eml",
            b"Subject: Vega\nMIME-Version: 1.0\nContent-Type: text/plain; charset=utf-8\n\nVega ricambi",
            "email_headers",
        ),
        (
            "data.mhtml",
            b"MIME-Version: 1.0\nContent-Type: text/html; charset=utf-8\n\n<html><body>Vega ricambi</body></html>",
            "document_text",
        ),
        ("data.xls", xls.getvalue(), "dataset_row"),
        ("data.ods", ods.getvalue(), "dataset_row"),
        ("data.pptx", pptx.getvalue(), "document_text"),
        (
            "data.html",
            b"<html><body><h1>Vega ricambi</h1><p>Priorita P1</p></body></html>",
            "document_text",
        ),
        ("data.sql", "CREATE TABLE città(id INT);".encode("utf-16"), "ddl_table"),
    ]


@pytest.mark.parametrize(
    "name,data,kind", fixtures(), ids=lambda v: v if isinstance(v, str) else None
)
def test_format_worker_real(tmp_path, name, data, kind):
    app = Application(tmp_path)
    rev = app.ingest(name, data)["revision_id"]
    result = app.parse(rev)
    assert result["status"] == "success", result
    assert any(e["type"] == kind for e in app.evidence())
    if name.endswith(".ods"):
        assert any("of:=1+2" in e["text"] for e in app.evidence())
    if name.endswith(".yaml"):
        assert any(
            e["data"].get("value", {}).get("date") == "2026-01-01"
            for e in app.evidence()
        )
